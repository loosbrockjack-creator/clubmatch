#!/usr/bin/env python3
"""Build active/leads/clubmatch_exec_contacts.xlsx from the normalized contacts.

Tabs: Contacts | Clubs | Needs Review | Sources & Notes
Designed to be shared with group members and to work for manual sending,
mail merge, or CRM import.
"""
import json, pathlib, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = pathlib.Path(__file__).resolve().parents[2]
SRC = ROOT / "research/outreach/contacts_normalized.json"
OUT = ROOT / "leads/clubmatch_exec_contacts.xlsx"

FONT = "Arial"
INK = "1F2937"        # charcoal header
ACCENT = "3B5BA5"     # restrained blue for section labels
MUTED = "6B7280"
BAND = "F3F4F6"
WARN = "FFF4E5"

H_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
H_FILL = PatternFill("solid", fgColor=INK)
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color="1155CC", underline="single")
MUTED_F = Font(name=FONT, size=9, color=MUTED)
TITLE_F = Font(name=FONT, size=14, bold=True, color=INK)
SECTION_F = Font(name=FONT, size=11, bold=True, color=ACCENT)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(bottom=THIN)


def header(ws, cols, row=1):
    for i, (title, width) in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=title)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[row].height = 30


def finish(ws, ncols, nrows, hdr_row=1, name=None):
    """Freeze header, add autofilter, band rows."""
    ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)
    if nrows:
        ws.auto_filter.ref = (f"A{hdr_row}:{get_column_letter(ncols)}{hdr_row + nrows}")
    for r in range(hdr_row + 1, hdr_row + nrows + 1):
        if (r - hdr_row) % 2 == 0:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=BAND)
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).border = BORDER


def put(ws, r, c, value, *, font=None, link=None, wrap=False, center=False):
    cell = ws.cell(row=r, column=c, value=value)
    cell.font = font or BODY
    if link:
        cell.hyperlink = link
        cell.font = LINK
    cell.alignment = Alignment(vertical="top", wrap_text=wrap,
                               horizontal="center" if center else None)
    return cell


def main():
    data = json.loads(SRC.read_text())
    contacts, review, clubs = data["contacts"], data["needs_review"], data["clubs"]
    today = data.get("generated") or datetime.date.today().isoformat()

    wb = Workbook()

    # ---------------------------------------------------------------- Contacts
    ws = wb.active
    ws.title = "Contacts"
    cols = [("First Name", 14), ("Last Name", 16), ("Role", 26), ("Email", 28),
            ("Club", 40), ("College", 34), ("Person Type", 15), ("Key Officer", 11),
            ("Major / Department", 26), ("Year / Title", 20), ("Phone", 15),
            ("Also In Other Clubs", 28), ("Source (StuOrg)", 18),
            ("Verified In ISU Directory", 13), ("Last Verified", 13), ("Confidence", 11)]
    header(ws, cols)
    KEY = {"president", "vice president", "treasurer", "secretary"}
    for i, r in enumerate(contacts, start=2):
        is_key = any(p.strip().lower() in KEY for p in r["role"].split("/")) or r["person_type"] == "Advisor"
        put(ws, i, 1, r["first_name"]); put(ws, i, 2, r["last_name"])
        put(ws, i, 3, r["role"], wrap=True)
        put(ws, i, 4, r["email"], link=f"mailto:{r['email']}")
        put(ws, i, 5, r["club"], wrap=True); put(ws, i, 6, r["college"], wrap=True)
        put(ws, i, 7, r["person_type"]); put(ws, i, 8, "Yes" if is_key else "No", center=True)
        put(ws, i, 9, r["major_or_department"], wrap=True)
        put(ws, i, 10, r["classification_or_title"], wrap=True)
        put(ws, i, 11, r["phone"])
        put(ws, i, 12, r["other_clubs"], wrap=True)
        put(ws, i, 13, "Officers page", link=r["source_url"])
        put(ws, i, 14, r["directory_verified"], center=True)
        put(ws, i, 15, r["last_verified"], center=True)
        put(ws, i, 16, r["confidence"], center=True)
    finish(ws, len(cols), len(contacts))

    # ------------------------------------------------------------------- Clubs
    ws2 = wb.create_sheet("Clubs")
    cols2 = [("Club", 40), ("College", 34), ("Members", 10), ("Verified Contacts", 12),
             ("Needs Review", 11), ("Coverage", 18), ("Club Contact Email", 28),
             ("Role-Based Inbox (durable)", 28), ("Instagram", 34), ("LinkedIn", 30),
             ("Discord", 26), ("Facebook", 32), ("X / Twitter", 26), ("YouTube", 28),
             ("Club Website", 26), ("Other Link", 30), ("Other Link Type", 24),
             ("StuOrg Page", 16), ("Officers Page", 16), ("Last Verified", 13)]
    header(ws2, cols2)
    for i, c in enumerate(clubs, start=2):
        put(ws2, i, 1, c["club"], wrap=True); put(ws2, i, 2, c["college"], wrap=True)
        put(ws2, i, 3, c["members"], center=True)
        # live counts so the sheet stays correct if Contacts is filtered or edited
        ws2.cell(row=i, column=4, value=f'=COUNTIF(Contacts!$E:$E,$A{i})').font = BODY
        ws2.cell(row=i, column=4).alignment = Alignment(horizontal="center")
        ws2.cell(row=i, column=5, value=f"=COUNTIF('Needs Review'!$E:$E,$A{i})").font = BODY
        ws2.cell(row=i, column=5).alignment = Alignment(horizontal="center")
        put(ws2, i, 6, c["coverage"])
        if c["listed_contact_email"]:
            put(ws2, i, 7, c["listed_contact_email"], link=f"mailto:{c['listed_contact_email']}")
        if c["role_based_inbox"]:
            put(ws2, i, 8, c["role_based_inbox"], link=f"mailto:{c['role_based_inbox']}")
        for col, key in ((9, "instagram"), (10, "linkedin"), (11, "discord"),
                         (12, "facebook"), (13, "x_twitter"), (14, "youtube"), (15, "club_website"),
                         (16, "other_link")):
            if c[key]:
                put(ws2, i, col, c[key], link=c[key], wrap=True)
        put(ws2, i, 17, c["other_link_type"], wrap=True)
        put(ws2, i, 18, "StuOrg", link=c["stuorg_url"])
        put(ws2, i, 19, "Officers", link=c["officers_url"])
        put(ws2, i, 20, c["last_verified"], center=True)
    finish(ws2, len(cols2), len(clubs))

    # ------------------------------------------------------------ Needs Review
    ws3 = wb.create_sheet("Needs Review")
    note = ws3.cell(row=1, column=1,
                    value="Listed on the club's official StuOrg officers page, but no matching current record in the ISU "
                          "directory, most likely graduated or directory-suppressed. Not included in Contacts. "
                          "Verify before using.")
    note.font = Font(name=FONT, size=10, bold=True, color="92400E")
    note.alignment = Alignment(vertical="center", wrap_text=True)
    ws3.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
    ws3.row_dimensions[1].height = 30
    for c in range(1, 10):
        ws3.cell(row=1, column=c).fill = PatternFill("solid", fgColor=WARN)
    cols3 = [("First Name", 14), ("Last Name", 16), ("Role", 26), ("Email", 28),
             ("Club", 40), ("College", 30), ("Why It Needs Review", 60),
             ("Source (StuOrg)", 18), ("Last Checked", 13)]
    header(ws3, cols3, row=2)
    for i, r in enumerate(review, start=3):
        put(ws3, i, 1, r["first_name"]); put(ws3, i, 2, r["last_name"])
        put(ws3, i, 3, r["role"], wrap=True)
        put(ws3, i, 4, r["email"], link=f"mailto:{r['email']}")
        put(ws3, i, 5, r["club"], wrap=True); put(ws3, i, 6, r["college"], wrap=True)
        put(ws3, i, 7, r["review_reason"], wrap=True)
        put(ws3, i, 8, "Officers page", link=r["source_url"])
        put(ws3, i, 9, r["last_verified"], center=True)
    finish(ws3, len(cols3), len(review), hdr_row=2)

    # -------------------------------------------------------- Sources & Notes
    ws4 = wb.create_sheet("Sources & Notes")
    ws4.column_dimensions["A"].width = 34
    ws4.column_dimensions["B"].width = 96
    r = 1
    ws4.cell(row=r, column=1, value="ClubMatch | Iowa State Exec Board Contact List").font = TITLE_F
    r += 1
    ws4.cell(row=r, column=1, value=f"Built {today}  |  Pilot set: 25 academic and professional clubs").font = MUTED_F
    r += 2

    def section(label):
        nonlocal r
        c = ws4.cell(row=r, column=1, value=label)
        c.font = SECTION_F
        r += 1

    def kv(k, v, *, formula=False):
        nonlocal r
        a = ws4.cell(row=r, column=1, value=k); a.font = Font(name=FONT, size=10, bold=True)
        a.alignment = Alignment(vertical="top")
        b = ws4.cell(row=r, column=2, value=v); b.font = BODY
        b.alignment = Alignment(vertical="top", wrap_text=True)
        r += 1

    section("Summary")
    kv("Clubs in this pilot", f"=COUNTA(Clubs!$A$2:$A${1+len(clubs)})")
    kv("Verified contacts", f"=COUNTA(Contacts!$D$2:$D${1+len(contacts)})")
    kv("Rows needing review", f"=COUNTA('Needs Review'!$D$3:$D${2+len(review)})")
    kv("Clubs with at least one verified contact",
       f'=SUMPRODUCT(--(COUNTIF(Contacts!$E$2:$E${1+len(contacts)},Clubs!$A$2:$A${1+len(clubs)})>0))')
    kv("Advisors included (durable contacts)",
       f'=COUNTIF(Contacts!$G$2:$G${1+len(contacts)},"Advisor")')
    r += 1

    section("Where this came from")
    kv("Primary source", "Iowa State University Student Organization Database (StuOrg), each club's public "
                         "Officers page: https://www.stuorg.iastate.edu/<club>/officers")
    kv("How emails were obtained", "On the StuOrg officers page each officer's name links to the ISU directory with "
                                   "their ISU address embedded in the link. Those addresses are published by the clubs "
                                   "themselves on ISU-hosted public pages. No sign-in was used.")
    kv("Verification", "Every address was looked up in the public ISU directory (info.iastate.edu). A row is marked "
                       "Verified only when the directory returns a person whose first and last name match the name the "
                       "club listed. Majors, class years, departments, titles and office phone numbers come from that "
                       "same directory record.")
    kv("Iowa State only", "Every contact in this workbook is an @iastate.edu address confirmed against the Iowa State "
                          "directory. No contacts from any other university are included.")
    kv("Collection method", "Read-only public page requests with a polite delay. Nothing was submitted, joined, or "
                            "messaged, and no login was used.")
    r += 1

    section("How to use it")
    kv("Mail merge / CRM import", "The Contacts tab is one row per person with First Name, Last Name, Email, Role, Club "
                                  "and College as separate columns, so it imports directly.")
    kv("Who to contact first", 'Filter Contacts on Key Officer = "Yes" for presidents, vice presidents, treasurers, '
                               "secretaries and advisors. Those are the people who can actually say yes.")
    kv("Avoid double sending", "A few people sit on two different exec boards. The Also In Other Clubs column flags "
                               "them, so de-duplicate on Email if you send one message per person.")
    kv("Link columns on the Clubs tab", "Club Website is the club's own site. Other Link is something else the club "
                                        "lists on StuOrg, with Other Link Type saying what it is: a national parent "
                                        "organization or a signup form. National-org links are reference only, not "
                                        "Iowa State contacts, so do not use them for outreach. Links hosted at another "
                                        "university and links that were malformed were dropped.")
    kv("Most durable channels", "Role-based club inboxes on the Clubs tab (the shared club addresses, not a "
                                "person's netid) and "
                                "advisors outlast student turnover. Prefer them for anything you want to still work "
                                "next semester.")
    r += 1

    section("Limits and cautions")
    kv("Officer turnover", "Exec boards turn over every spring. Student rows go stale within a semester or two. "
                           "Re-run the scripts in active/execution/scripts/ to refresh.")
    kv("Staleness is already visible", "The Needs Review tab holds people a club still lists who no longer appear in "
                                       "the ISU directory. That is the clearest signal a club's page is out of date.")
    kv("Student privacy", "These are student email addresses. Keep this file inside the group, do not republish it, "
                          "send only relevant outreach, and include a clear opt-out in every message. Honor any "
                          "request to be removed immediately.")
    kv("Bulk email rules", "Follow CAN-SPAM and Iowa State's policies on contacting student organizations. If the "
                           "group plans a large send, the Student Activities Center or club advisors are a safer and "
                           "more credible intro path than cold-emailing officers.")
    kv("Not an ISU endorsement", "StuOrg pages carry Iowa State's disclaimer that club page contents are not reviewed "
                                 "or approved by the university.")
    kv("Classification note", "College labels come from the project's existing club classifier. A few are imperfect "
                              "(for example the SAE student chapter is grouped under Ivy College of Business). Treated "
                              "as given here, not corrected.")
    r += 1

    section("Column definitions, Contacts tab")
    for k, v in [("Role", "Exactly as the club lists it. A single person holding two roles in one club shows both, "
                          "separated by a slash."),
                 ("Person Type", "Student Officer or Advisor, from the ISU directory record."),
                 ("Key Officer", "Yes for president, vice president, treasurer, secretary or advisor."),
                 ("Major / Department", "Major for students, department for faculty and staff."),
                 ("Year / Title", "Class year for students, job title for faculty and staff."),
                 ("Phone", "Office phone where the directory publishes one. Student phone numbers are not public, so "
                           "this is populated mainly for advisors."),
                 ("Verified In ISU Directory", "Yes means the address resolves to an ISU person whose name matches the "
                                               "club's listing."),
                 ("Confidence", "High means name and address both confirmed in the ISU directory.")]:
        kv(k, v)
    r += 1

    section("Rebuild")
    kv("1. Scrape officers", "python3 active/execution/scripts/scrape_officers.py")
    kv("2. Verify against directory", "python3 active/execution/scripts/verify_directory.py")
    kv("3. Normalize and join", "python3 active/execution/scripts/normalize_contacts.py")
    kv("4. Build this workbook", "python3 active/execution/scripts/build_workbook.py")
    kv("Club input list", "active/research/outreach/pilot_clubs.json")
    kv("Not committed to git", "active/leads/ and active/research/outreach/ are gitignored because this file contains "
                               "personal contact data.")

    for row in range(1, r + 1):
        ws4.row_dimensions[row].height = None
    ws4.sheet_view.showGridLines = False

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"contacts={len(contacts)} clubs={len(clubs)} needs_review={len(review)}")
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
